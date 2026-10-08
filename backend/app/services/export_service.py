import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from typing import List, Dict, Any, Tuple
from sqlalchemy.orm import Session

from backend.app.models.clause import ClauseModel
from backend.app.models.document import DocumentModel
from backend.app.models.correction import CorrectionModel

class ExportService:
    DISCLAIMER_TEXT = "INFORMATIONAL NOTICE: This document was processed by LexClear. It is provided for informational and drafting assistance, not formal legal advice."

    @classmethod
    def export_docx(cls, document_id: str, mode: str, db: Session) -> io.BytesIO:
        doc_record = db.query(DocumentModel).filter(DocumentModel.id == document_id).first()
        clauses = db.query(ClauseModel).filter(ClauseModel.document_id == document_id).order_by(ClauseModel.clause_index).all()
        corrections = db.query(CorrectionModel).filter(CorrectionModel.document_id == document_id).all()
        corr_map = {c.clause_id: c for c in corrections}

        doc = Document()
        
        # Header disclaimer
        disclaimer_p = doc.add_paragraph()
        run = disclaimer_p.add_run(cls.DISCLAIMER_TEXT + "\n")
        run.font.size = Pt(8.5)
        run.font.italic = True
        run.font.color.rgb = RGBColor(120, 120, 120)

        # Title
        title_p = doc.add_paragraph()
        t_run = title_p.add_run(f"LEGAL DOCUMENT: {doc_record.file_name if doc_record else 'Agreement'}\n")
        t_run.font.size = Pt(16)
        t_run.font.bold = True

        for c in clauses:
            # Heading
            h_p = doc.add_paragraph()
            h_run = h_p.add_run(c.title)
            h_run.font.size = Pt(12)
            h_run.font.bold = True

            p = doc.add_paragraph()
            corr = corr_map.get(c.id)

            if mode == "tracked_changes" and corr and corr.revised_text != c.text:
                # Show strikethrough for original, green for revised
                del_run = p.add_run(f"[REVISED DELETION: {c.text}] ")
                del_run.font.strike = True
                del_run.font.color.rgb = RGBColor(180, 40, 40)

                ins_run = p.add_run(f"\n[INSERTION: {corr.revised_text}]\n")
                ins_run.font.bold = True
                ins_run.font.color.rgb = RGBColor(30, 130, 40)
            elif corr and corr.status == "accepted":
                p.add_run(corr.revised_text)
            else:
                p.add_run(c.text)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer

    @classmethod
    def export_pdf(cls, document_id: str, db: Session) -> io.BytesIO:
        doc_record = db.query(DocumentModel).filter(DocumentModel.id == document_id).first()
        clauses = db.query(ClauseModel).filter(ClauseModel.document_id == document_id).order_by(ClauseModel.clause_index).all()
        corrections = db.query(CorrectionModel).filter(CorrectionModel.document_id == document_id).all()
        corr_map = {c.clause_id: c for c in corrections}

        buffer = io.BytesIO()
        pdf = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=54, rightMargin=54, topMargin=54, bottomMargin=54)
        
        styles = getSampleStyleSheet()
        disclaimer_style = ParagraphStyle(
            'Disclaimer',
            parent=styles['Normal'],
            fontSize=8,
            leading=10,
            textColor='#666666',
            alignment=1
        )
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontSize=16,
            leading=20,
            textColor='#1a202c',
            spaceAfter=15
        )
        clause_title_style = ParagraphStyle(
            'ClauseTitle',
            parent=styles['Heading2'],
            fontSize=11,
            leading=14,
            textColor='#2b6cb0',
            spaceBefore=10,
            spaceAfter=4
        )
        body_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor='#2d3748',
            spaceAfter=8
        )

        story = []
        story.append(Paragraph(cls.DISCLAIMER_TEXT, disclaimer_style))
        story.append(Spacer(1, 15))
        story.append(Paragraph(f"<b>{doc_record.file_name if doc_record else 'Legal Agreement'}</b>", title_style))

        for c in clauses:
            story.append(Paragraph(f"<b>{c.title}</b>", clause_title_style))
            corr = corr_map.get(c.id)
            text_to_show = corr.revised_text if (corr and corr.status == "accepted") else c.text
            # Clean newlines for ReportLab Paragraph
            formatted_text = text_to_show.replace("\n", "<br/>")
            story.append(Paragraph(formatted_text, body_style))
            story.append(Spacer(1, 4))

        pdf.build(story)
        buffer.seek(0)
        return buffer
