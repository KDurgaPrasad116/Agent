"""
<<<<<<< HEAD
Unified file parser for various formats.
=======
Unified file parser for PDF, DOCX, and TXT formats.
Replaces the original services/input_processor.py logic
but is importable from both old and new code paths.
>>>>>>> 3a27112 (Pushing major agent changes)
"""

import io
from PyPDF2 import PdfReader
from docx import Document
<<<<<<< HEAD
import pandas as pd
from bs4 import BeautifulSoup

# Lazy load easyocr since it can be heavy
_reader = None
def get_ocr_reader():
    global _reader
    if _reader is None:
        import easyocr
        _reader = easyocr.Reader(['en'])
    return _reader
=======

>>>>>>> 3a27112 (Pushing major agent changes)

def parse_file(file) -> str:
    """
    Extract plain text from an uploaded FastAPI UploadFile.
<<<<<<< HEAD
    """
    content = file.file.read()
    return parse_bytes(content, file.filename)
=======
    Supports: .pdf, .docx, .txt
    """
    filename = file.filename.lower()

    if filename.endswith(".pdf"):
        return _extract_pdf(file)
    elif filename.endswith(".docx"):
        return _extract_docx(file)
    elif filename.endswith(".txt"):
        return _extract_txt(file)
    else:
        raise ValueError(
            f"Unsupported file format '{filename}'. "
            "Accepted formats: .pdf, .docx, .txt"
        )

>>>>>>> 3a27112 (Pushing major agent changes)

def parse_bytes(content: bytes, filename: str) -> str:
    """
    Extract plain text from raw bytes.
    Useful when the file has already been read into memory.
    """
    filename = filename.lower()

    if filename.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(content))
        return "".join(p.extract_text() or "" for p in reader.pages)

    elif filename.endswith(".docx"):
        doc = Document(io.BytesIO(content))
        return "\n".join(p.text for p in doc.paragraphs)

<<<<<<< HEAD
    elif filename.endswith((".txt", ".md", ".json")):
=======
    elif filename.endswith(".txt"):
>>>>>>> 3a27112 (Pushing major agent changes)
        try:
            return content.decode("utf-8")
        except UnicodeDecodeError:
            return content.decode("latin-1")

<<<<<<< HEAD
    elif filename.endswith((".csv", ".xlsx", ".xls")):
        if filename.endswith(".csv"):
            df = pd.read_csv(io.BytesIO(content))
        else:
            df = pd.read_excel(io.BytesIO(content))
        return df.to_string(index=False)

    elif filename.endswith(".html"):
        soup = BeautifulSoup(content, "html.parser")
        return soup.get_text(separator="\n")

    elif filename.endswith((".png", ".jpg", ".jpeg")):
        reader = get_ocr_reader()
        result = reader.readtext(content, detail=0)
        return "\n".join(result)

    else:
        raise ValueError(
            f"Unsupported file format '{filename}'. "
            "Accepted formats: .pdf, .docx, .txt, .md, .json, .csv, .xlsx, .xls, .html, .png, .jpg, .jpeg"
        )
=======
    else:
        raise ValueError(f"Unsupported format: {filename}")


# ── Private helpers ───────────────────────────────────────────────────────────

def _extract_pdf(file) -> str:
    reader = PdfReader(file.file)
    return "".join(p.extract_text() or "" for p in reader.pages)


def _extract_docx(file) -> str:
    doc = Document(file.file)
    return "\n".join(p.text for p in doc.paragraphs)


def _extract_txt(file) -> str:
    content = file.file.read()
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("latin-1")
>>>>>>> 3a27112 (Pushing major agent changes)
