"""
Unified file parser for PDF, DOCX, and TXT formats.
Replaces the original services/input_processor.py logic
but is importable from both old and new code paths.
"""

import io
from PyPDF2 import PdfReader
from docx import Document


def parse_file(file) -> str:
    """
    Extract plain text from an uploaded FastAPI UploadFile.
    Supports: .pdf, .docx, .txt
    """
    filename = file.filename.lower()

    if filename.endswith(".pdf"):
        return _extract_pdf(file)
    elif filename.endswith(".docx"):
        return _extract_docx(file)
    elif filename.endswith(".txt"):
        return _extract_txt(file)
    else:
        raise ValueError(
            f"Unsupported file format '{filename}'. "
            "Accepted formats: .pdf, .docx, .txt"
        )


def parse_bytes(content: bytes, filename: str) -> str:
    """
    Extract plain text from raw bytes.
    Useful when the file has already been read into memory.
    """
    filename = filename.lower()

    if filename.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(content))
        return "".join(p.extract_text() or "" for p in reader.pages)

    elif filename.endswith(".docx"):
        doc = Document(io.BytesIO(content))
        return "\n".join(p.text for p in doc.paragraphs)

    elif filename.endswith(".txt"):
        try:
            return content.decode("utf-8")
        except UnicodeDecodeError:
            return content.decode("latin-1")

    else:
        raise ValueError(f"Unsupported format: {filename}")


# ── Private helpers ───────────────────────────────────────────────────────────

def _extract_pdf(file) -> str:
    reader = PdfReader(file.file)
    return "".join(p.extract_text() or "" for p in reader.pages)


def _extract_docx(file) -> str:
    doc = Document(file.file)
    return "\n".join(p.text for p in doc.paragraphs)


def _extract_txt(file) -> str:
    content = file.file.read()
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("latin-1")
